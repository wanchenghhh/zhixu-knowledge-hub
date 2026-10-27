"""Export generated Markdown with headings, lists and tables to a Word document."""

from io import BytesIO
import re
from docx import Document


def inline(paragraph, text):
    for i, part in enumerate(re.split(r"(\*\*.*?\*\*)", text)):
        run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
        if part.startswith("**") and part.endswith("**"):
            run.bold = True


def word_document(markdown):
    document = Document()
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if (
            "|" in line
            and index + 1 < len(lines)
            and re.fullmatch(r"[| :\-]+", lines[index + 1].strip())
            and "-" in lines[index + 1]
        ):
            header = [s.strip() for s in line.strip("|").split("|")]
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, header, strict=True):
                inline(cell.paragraphs[0], text)
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                values = [s.strip() for s in lines[index].strip().strip("|").split("|")]
                cells = table.add_row().cells
                for i, text in enumerate(values[: len(cells)]):
                    inline(cells[i].paragraphs[0], text)
                index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)", line)
        if heading:
            document.add_heading(heading[2], level=min(3, len(heading[1])))
        elif re.match(r"^[-*]\s+", line):
            inline(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+[.)]\s+", line):
            inline(document.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", line))
        else:
            inline(document.add_paragraph(), line)
        index += 1
    output = BytesIO()
    document.save(output)
    return output.getvalue()
