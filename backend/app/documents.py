import csv
import hashlib
import json
import time
import zipfile
from pathlib import Path
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.orm import Session
from .config import decrypt, settings
from .db import File, Job, Setting, utcnow

EXTENSIONS = {"pdf", "doc", "docx", "md", "png", "jpg", "jpeg", "xlsx", "xls", "csv"}


def file_path(file: File):
    return settings().data_dir / "files" / file.id


def zip_guard(path):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 5000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
            raise ValueError("文档解压后过大或包含过多文件")


def block(text, location, **metadata):
    return {"text": text.strip(), "location": location, "metadata": metadata}


def table_blocks(rows, location, **metadata):
    cleaned = [
        [str(c if c is not None else "").replace("\n", " ").replace("|", "\\|") for c in row] for row in rows
    ]
    if not cleaned:
        return []
    header = " | ".join(cleaned[0])
    result = []
    for start in range(1, len(cleaned), 8):
        text = (
            header
            + "\n"
            + " | ".join(["---"] * len(cleaned[0]))
            + "\n"
            + "\n".join(" | ".join(row) for row in cleaned[start : start + 8])
        )
        result.append(
            block(
                text,
                f"{location} · 第 {start + 1}–{min(len(cleaned), start + 8)} 行",
                kind="table",
                **metadata,
            )
        )
    return result or [block(header, location, kind="table", **metadata)]


def cloud_call(method, *args):
    try:
        return method(*args)
    except Exception:
        raise ValueError(
            "云端文档解析请求失败，请核对 DocMind 凭证、权限、余额与网络；未自动重复提交"
        ) from None


def cloud_parse(path: Path, filename: str, file: File, db: Session, job: Job, cache_key="document"):
    """Persist cloud job IDs before polling; uncertain submissions never auto-repeat."""
    from alibabacloud_docmind_api20220711.client import Client
    from alibabacloud_docmind_api20220711 import models
    from alibabacloud_tea_openapi.models import Config
    from alibabacloud_tea_util.models import RuntimeOptions

    config_row = db.get(Setting, "ocr")
    config = config_row.value if config_row else {}
    if not config.get("enabled") or not config.get("id_secret") or not config.get("key_secret"):
        raise ValueError("此文件需要云端 OCR，请先配置并启用阿里云 DocMind")
    client = Client(
        Config(
            access_key_id=decrypt(config["id_secret"]),
            access_key_secret=decrypt(config["key_secret"]),
            endpoint="docmind-api.cn-hangzhou.aliyuncs.com",
            connect_timeout=15000,
            read_timeout=60000,
        )
    )
    previous = file.cloud_jobs.get(cache_key, {})
    cloud_id = previous.get("id")
    if previous.get("submitting") and not cloud_id:
        raise ValueError(
            "OCR 提交结果未知，为避免重复计费未自动重提；请在云端核对后删除此失败文件，再重新上传"
        )
    if not cloud_id:
        file.cloud_jobs = {**file.cloud_jobs, cache_key: {"submitting": True}}
        db.commit()
        with path.open("rb") as stream:
            response = cloud_call(
                client.submit_doc_structure_job_advance,
                models.SubmitDocStructureJobAdvanceRequest(file_url_object=stream, file_name=filename),
                RuntimeOptions(autoretry=False),
            )
        data = response.body.to_map().get("Data", {})
        cloud_id = data.get("Id")
        if not cloud_id:
            raise ValueError("OCR 未返回任务编号，请核对云端任务后重试")
        file.cloud_jobs = {**file.cloud_jobs, cache_key: {"id": cloud_id}}
        db.commit()
    deadline = time.monotonic() + 7200
    while time.monotonic() < deadline:
        job.touched = utcnow()
        db.commit()
        response = cloud_call(
            client.get_doc_structure_result,
            models.GetDocStructureResultRequest(id=cloud_id, reveal_markdown=True),
        )
        body = response.body.to_map()
        data = body.get("Data") or {}
        data = json.loads(data) if isinstance(data, str) else data
        status = body.get("Status", data.get("Status", ""))
        if status in {"Fail", "Failed", "Failure", "Timeout"}:
            raise ValueError(f"云端 OCR 任务失败：{status}；任务号 {cloud_id}")
        raw = data.get("DocStructureResult") or (data if data.get("layouts") else None)
        if raw:
            parsed = json.loads(raw) if isinstance(raw, str) else raw
            results = []
            layouts = parsed.get("layouts", [])
            if not layouts:
                for page in parsed.get("pages", []):
                    layouts.extend(page.get("layouts", []))
            for item in layouts:
                pages = item.get("pageNum", [0])
                if isinstance(pages, int):
                    pages = [pages]
                pages = [int(p) + 1 for p in pages]
                text = item.get("markdownContent") or item.get("text") or ""
                if not text and item.get("cells"):
                    cells = json.loads(item["cells"]) if isinstance(item["cells"], str) else item["cells"]
                    text = " | ".join(str(cell.get("text", "")) for cell in cells)
                if text.strip():
                    results.append(
                        block(
                            text,
                            f"第 {'–'.join(map(str, pages))} 页",
                            page=pages[0],
                            pages=pages,
                            coordinates=item.get("pos"),
                            kind=item.get("type", "text"),
                        )
                    )
            if not results:
                raise ValueError("OCR 返回结果未包含可用版面文字，请检查接口输出结构")
            return results
        if status in {"Success", "Succeeded"}:
            raise ValueError("OCR 已完成但未返回结构化内容")
        time.sleep(10)
    raise ValueError(f"OCR 查询超时，保留云端任务号供后续恢复：{cloud_id}")


def parse_file(file: File, db: Session, job: Job):
    path, ext = file_path(file), file.extension
    if ext == "pdf":
        import pdfplumber

        with pdfplumber.open(path) as document:
            if len(document.pages) > settings().max_pdf_pages:
                raise ValueError(f"PDF 超过 {settings().max_pdf_pages} 页，请拆分后上传")
    if ext in {"docx", "xlsx"}:
        zip_guard(path)
    if ext in {"png", "jpg", "jpeg", "doc"} or (
        ext in {"pdf", "docx"} and file.options.get("parseMode") == "enhanced"
    ):
        return cloud_parse(path, file.name, file, db, job)
    if ext == "md":
        text = path.read_bytes().decode("utf-8-sig")
        blocks, heading, start, lines = [], "文档", 1, []
        for number, line in enumerate(text.splitlines(), 1):
            if line.startswith("#") and lines:
                blocks.append(
                    block(
                        "\n".join(lines),
                        f"{heading} · 第 {start}–{number - 1} 行",
                        lineStart=start,
                        lineEnd=number - 1,
                    )
                )
                lines, start = [], number
            if line.startswith("#"):
                heading = line.lstrip("# ").strip()[:150]
            lines.append(line)
        if lines:
            blocks.append(
                block(
                    "\n".join(lines), f"{heading} · 第 {start}–{len(text.splitlines())} 行", lineStart=start
                )
            )
        return blocks
    if ext == "csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return table_blocks(list(csv.reader(stream)), "CSV")
    if ext in {"xls", "xlsx"}:
        result = []
        if ext == "xlsx":
            from openpyxl import load_workbook

            book = load_workbook(path.open("rb"), read_only=True, data_only=False)
            try:
                for sheet in book:
                    if sheet.max_row > 100000:
                        raise ValueError("工作表超过 100,000 行，请拆分后上传")
                    result.extend(
                        table_blocks(list(sheet.iter_rows(values_only=True)), sheet.title, sheet=sheet.title)
                    )
            finally:
                book.close()
        else:
            import xlrd

            book = xlrd.open_workbook(path)
            for sheet in book.sheets():
                if sheet.nrows > 100000:
                    raise ValueError("工作表超过 100,000 行")
                result.extend(
                    table_blocks(
                        [sheet.row_values(i) for i in range(sheet.nrows)], sheet.name, sheet=sheet.name
                    )
                )
        return result
    if ext == "docx":
        from docx import Document as Word

        doc = Word(path)
        result = [
            block(p.text, f"第 {i + 1} 段", paragraph=i + 1)
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip()
        ]
        for i, table in enumerate(doc.tables):
            result.extend(
                table_blocks([[cell.text for cell in row.cells] for row in table.rows], f"第 {i + 1} 个表格")
            )
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith("word/media/") and name.lower().endswith((".png", ".jpg", ".jpeg")):
                    data = archive.read(name)
                    image_path = (
                        settings().data_dir
                        / "files"
                        / (file.id + "-" + hashlib.sha256(data).hexdigest()[:12])
                    )
                    image_path.write_bytes(data)
                    try:
                        image_blocks = cloud_parse(image_path, Path(name).name, file, db, job, cache_key=name)
                        result.extend(
                            {
                                **b,
                                "location": f"文档插图 {Path(name).name}",
                                "metadata": {**b["metadata"], "image": name},
                            }
                            for b in image_blocks
                        )
                    finally:
                        image_path.unlink(missing_ok=True)
        return result
    if ext == "pdf":
        import pdfplumber

        result = []
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > settings().max_pdf_pages:
                raise ValueError(f"PDF 超过 {settings().max_pdf_pages} 页，请拆分后上传")
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if len(text.strip()) < 20 or page.images:
                    return cloud_parse(path, file.name, file, db, job)
                tables = page.find_tables()

                def outside_tables(obj):
                    return not any(
                        t.bbox[0] <= obj.get("x0", -1) <= t.bbox[2]
                        and t.bbox[1] <= obj.get("top", -1) <= t.bbox[3]
                        for t in tables
                    )

                text = page.filter(outside_tables).extract_text() or ""
                if text.strip():
                    result.append(block(text, f"第 {i + 1} 页", page=i + 1))
                for j, table in enumerate(tables):
                    result.extend(table_blocks(table.extract(), f"第 {i + 1} 页 · 表格 {j + 1}", page=i + 1))
        return result
    raise ValueError("不支持的文档格式")


def split_blocks(blocks: list, options: dict) -> list[Document]:
    size, overlap = options.get("chunkSize", 800), options.get("overlap", 100)
    if overlap >= size:
        raise ValueError("重叠长度必须小于分块长度")
    if sum(len(item["text"]) for item in blocks) > settings().max_text_chars:
        raise ValueError("文档文字量过大，请拆分后重新上传")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", "。", "；", "，", " ", ""],
        add_start_index=True,
    )
    documents = [
        Document(
            page_content=item["text"],
            metadata={"location": item["location"], "block_index": index, **item.get("metadata", {})},
        )
        for index, item in enumerate(blocks)
        if item["text"].strip()
    ]
    return splitter.split_documents(documents)
