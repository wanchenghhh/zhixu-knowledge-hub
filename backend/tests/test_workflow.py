import io
import json
from uuid import uuid4
import httpx
from docx import Document
from openpyxl import Workbook
from sqlalchemy import select
from app.config import decrypt
from app.db import Chunk, File, Job, Model, session_factory
from app.jobs import run_job


def upload(
    client, kb, data=b"# Travel\nReimbursement requires invoices and payment receipts.\n", name="rules.md"
):
    response = client.post(
        f"/api/v1/knowledge/{kb['id']}/upload",
        files={"file": (name, data)},
        data={"consent": "true", "options": json.dumps({"chunkSize": 200, "overlap": 20, "extractMetadata": False})},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    run_job(result["task"]["id"])
    return result["file"]["id"]


def build(client, kb, file_id):
    response = client.post(f"/api/v1/knowledge/{kb['id']}/build", json={"fileIds": [file_id]})
    assert response.status_code == 200, response.text
    run_job(response.json()["id"])
    return response.json()["id"]


def test_model_secrets_and_real_capability_requests(client, models):
    state = client.get("/api/v1/state").json()
    assert "test-secret-never-expose" not in json.dumps(state)
    with session_factory()() as db:
        model = db.get(Model, models["Embedding"]["id"])
        assert model.secret != "test-secret-never-expose"
        assert decrypt(model.secret) == "test-secret-never-expose"
        assert model.status == "ready"


def test_complete_import_retrieval_chat_logs(client, models, knowledge):
    file_id = upload(client, knowledge, "# 费用报销\n报销需要发票与支付凭证，15 个工作日内提交。".encode())
    state = client.get("/api/v1/state").json()
    assert next(f for f in state["files"] if f["id"] == file_id)["status"] == "preview"
    job_id = build(client, knowledge, file_id)
    with session_factory()() as db:
        assert db.get(Job, job_id).status == "success"
        assert db.get(File, file_id).status == "ready"
    request = {"question": "报销需要什么材料？", "kbIds": [knowledge["id"]], "threshold": 0.3, "rerank": True}
    result = client.post("/api/v1/retrieval", json=request)
    assert result.status_code == 200, result.text
    assert "发票" in result.json()["sources"][0]["text"]
    response = client.post("/api/v1/chat", json={**request, "modelId": models["LLM"]["id"]})
    assert response.status_code == 200, response.text
    assert "event: token" in response.text and "event: done" in response.text
    state = client.get("/api/v1/state").json()
    log = state["logs"][0]
    assert log["status"] == "success", log
    assert log["tokens"] == 45, log
    assert log["sources"][0]["fileId"] == file_id
    assert "[999]" not in log["answer"] and "[来源不可用]" in log["answer"]
    assert state["sessions"][0]["messages"][-1]["sources"]
    assert client.get(f"/api/v1/files/{file_id}/download").content.startswith("# 费用报销".encode())


def test_roles_download_retrieval_and_browser_isolation(client, models, knowledge):
    file_id = upload(client, knowledge)
    build(client, knowledge, file_id)
    private = {**knowledge, "visible": ["admin"]}
    assert client.put(f"/api/v1/knowledge/{knowledge['id']}", json=private).status_code == 200
    headers = {"X-Role": "user"}
    assert client.get("/api/v1/state", headers=headers).json()["knowledge"] == []
    assert client.get(f"/api/v1/files/{file_id}/download", headers=headers).status_code == 404
    assert (
        client.post(
            "/api/v1/retrieval", headers=headers, json={"question": "x", "kbIds": [knowledge["id"]]}
        ).status_code
        == 404
    )
    assert client.delete(f"/api/v1/models/{models['LLM']['id']}", headers=headers).status_code == 403
    response = client.post(
        "/api/v1/chat", json={"question": "报销", "kbIds": [knowledge["id"]], "modelId": models["LLM"]["id"]}
    )
    assert response.status_code == 200
    state = client.get("/api/v1/state").json()
    conversation = state["sessions"][0]
    assert client.get("/api/v1/state", headers={"X-Session-ID": str(uuid4())}).json()["sessions"] == []
    assert (
        client.delete(
            f"/api/v1/conversations/{conversation['id']}", headers={"X-Session-ID": str(uuid4())}
        ).status_code
        == 404
    )


def test_disable_no_stale_retrieval_and_empty_answer(client, models, knowledge):
    file_id = upload(client, knowledge)
    build(client, knowledge, file_id)
    assert client.patch(f"/api/v1/files/{file_id}", json={"action": "disable"}).status_code == 200
    response = client.post(
        "/api/v1/chat", json={"question": "报销", "kbIds": [knowledge["id"]], "modelId": models["LLM"]["id"]}
    )
    assert "未找到足够依据" in response.text
    assert client.get("/api/v1/state").json()["logs"][0]["status"] == "empty"
    assert client.get(f"/api/v1/files/{file_id}/download", headers={"X-Role": "user"}).status_code == 404


def test_rebuild_failure_preserves_active_index(client, models, knowledge, cloud):
    file_id = upload(client, knowledge)
    build(client, knowledge, file_id)
    changed = {**knowledge, "embedding": models["Embedding"]["id"]}
    new = client.post(
        "/api/v1/models",
        json={
            "name": "new",
            "type": "Embedding",
            "model": "other",
            "provider": "AIHubMix",
            "url": "https://aihubmix.com/v1",
            "dimension": 64,
            "apiKey": "key",
        },
    ).json()
    assert client.post(f"/api/v1/models/{new['id']}/test").status_code == 200
    cloud.post("https://aihubmix.com/v1/embeddings").mock(return_value=httpx.Response(503))
    changed["embedding"] = new["id"]
    response = client.put(f"/api/v1/knowledge/{knowledge['id']}", json=changed)
    assert response.status_code == 200, response.text
    run_job(response.json()["task"]["id"])
    with session_factory()() as db:
        from app.db import Knowledge

        assert db.get(Knowledge, knowledge["id"]).embedding == models["Embedding"]["id"]
        assert db.scalar(select(Chunk).where(Chunk.file_id == file_id)) is not None


def test_word_excel_csv_parsers(client, knowledge):
    doc = Document()
    doc.add_paragraph("差旅管理条例要求员工提交合规发票。")
    buffer = io.BytesIO()
    doc.save(buffer)
    file_id = upload(client, knowledge, buffer.getvalue(), "rules.docx")
    book = Workbook()
    sheet = book.active
    sheet.title = "费用标准"
    sheet.append(["项目", "限额"])
    sheet.append(["住宿", 480])
    buffer = io.BytesIO()
    book.save(buffer)
    excel_id = upload(client, knowledge, buffer.getvalue(), "limits.xlsx")
    csv_id = upload(client, knowledge, "项目,限额\n餐费,80\n".encode(), "limits.csv")
    with session_factory()() as db:
        assert "第 1 段" in db.get(File, file_id).blocks[0]["location"]
        assert "费用标准" in db.get(File, excel_id).blocks[0]["location"]
        assert "480" in db.get(File, excel_id).blocks[0]["text"]
        assert "80" in db.get(File, csv_id).blocks[0]["text"]


def test_upload_consent_duplicate_and_validation(client, knowledge):
    response = client.post(f"/api/v1/knowledge/{knowledge['id']}/upload", files={"file": ("a.md", b"text")})
    assert response.status_code == 400
    upload(client, knowledge)
    response = client.post(
        f"/api/v1/knowledge/{knowledge['id']}/upload",
        files={"file": ("other.md", b"# Travel\nReimbursement requires invoices and payment receipts.\n")},
        data={"consent": "true", "options": json.dumps({"chunkSize": 200, "overlap": 20, "extractMetadata": False})},
    )
    assert response.status_code == 409
    assert client.post("/api/v1/models", json={"apiKey": "test-never-echo"}).status_code == 422
    assert "test-never-echo" not in client.post("/api/v1/models", json={"apiKey": "test-never-echo"}).text


def test_missing_ocr_fails_clearly_without_local_models(client, knowledge):
    file_id = upload(client, knowledge, b"image-content", "scan.png")
    with session_factory()() as db:
        file = db.get(File, file_id)
        assert file.status == "failed"
        assert "DocMind" in file.error


def test_no_unapproved_model_endpoints(client):
    response = client.post(
        "/api/v1/models",
        json={
            "name": "bad",
            "provider": "AIHubMix",
            "type": "LLM",
            "model": "x",
            "url": "http://127.0.0.1:22",
            "apiKey": "must-not-echo",
        },
    )
    assert response.status_code == 422
    assert "must-not-echo" not in response.text
